# coding=utf-8

from docx import Document
import xlwt
import os
import traceback

workbook = xlwt.Workbook()  # 新建一个工作簿
for group in os.listdir('.'):
    if group[-3:] == '.py'or group[-4:] == '.xls'or group=='.idea':
        continue
    '''
    --------------------------------------------
    0. 文件批量改名成 .docx格式
    --------------------------------------------
    '''
    from win32com import client as wc  # 导入模块

    path = group+"."  # 待处理文件夹
    #word = wc.Dispatch("word.application")  # 打开word应用程序（使用word打开）
    word = wc.Dispatch("Kwps.application")  # 打开word应用程序（使用wps打开）
    file_num_total=0
    for file in os.listdir(path):
        if file[-3:] == 'RTF' or file[-3:] == 'doc':
            file_num_total=file_num_total+1

    file_num_current=0
    if not os.path.exists(group+'\\'+'docx'):
        os.makedirs(group+'\\'+'docx')
    for file in os.listdir(path):
        if file[-3:] == 'RTF' or file[-3:] == 'doc':
            file_num_current = file_num_current + 1  # 当前处理的文件序号
            (file_path, temp_file_name) = os.path.split(file)
            (short_name, extension) = os.path.splitext(temp_file_name)  # short_name: 不带后缀的文件名

            print('转换文档 ' +group+'//'+file + '到docx文件格式..进度：' + str(file_num_current) + '/ ' + str(file_num_total))
            if not os.path.exists(os.getcwd()+'\\'+group+'\\docx\\' + short_name + ".docx"): # 如果docx已经存在就不再重新转换了
                doc = word.Documents.Open(os.getcwd()+'\\'+group+'\\'+file)
                doc.SaveAs(os.getcwd()+'\\'+group+'\\docx\\'+ short_name + ".docx", 12)  # 另存为后缀为".docx"的文件，其中参数0指doc文件
                doc.Close()
    word.Quit()

    '''
    --------------------------------------------
    1. 新建工作簿，写入标题栏
    --------------------------------------------
    '''
    sheet = workbook.add_sheet(group)  # 在工作簿中新建一个表格
    sheet_head='姓名,住院号,监测日期,性别,年龄,身高,体重,BMI,Date of birth,Date of study,总记录时间,' \
               '有效睡眠时间,总睡眠时间,睡眠潜伏期,睡眠效率,Wake时间,Wake占比,StageREM时间,StageREM占比,' \
               'Stage1时间,Stage1占比,Stage2时间,Stage2占比,Stage3时间,Stage3占比,Stage4时间,Stage4占比,' \
               '觉醒次数,觉醒指数,' \
               '阻塞性事件次数,阻塞性事件平均时间,阻塞性事件NREM下次数,阻塞性事件REM下次数,' \
               '混合性事件次数,混合性事件平均时间,混合性事件NREM下次数,混合性事件REM下次数,' \
               '中枢性事件次数,中枢性事件平均时间,中枢性事件NREM下次数,中枢性事件REM下次数,' \
               '所有暂停次数,所有暂停平均时间,所有暂停最长时间,所有暂停NREM下次数,所有暂停REM下次数,' \
               '低通气次数,低通气平均时间,低通气最长时间,低通气NREM下次数,低通气REM下次数,' \
               'AHI,OAI,MAD,鼾声次数,鼾声指数,' \
               '俯卧下阻塞性指数,左侧卧下阻塞性指数,右侧卧下阻塞性指数,仰卧下阻塞性指数,' \
               '俯卧下混合性指数,左侧卧下混合性指数,右侧卧下混合性指数,仰卧下混合性指数,' \
               '俯卧下中枢性指数,左侧卧下中枢性指数,右侧卧下中枢性指数,仰卧下中枢性指数,' \
               '俯卧下低通气指数,左侧卧下低通气指数,右侧卧下低通气指数,仰卧下低通气指数,' \
               '平均血氧,最低血氧,氧减值>=2%次数,氧减值>=3%次数,氧减值>=4%次数,ODI,氧减值>=5%次数,' \
               '血氧<95%时间（分钟）,血氧<90%时间（分钟）,血氧<85%时间（分钟）,血氧<80%时间（分钟）,' \
               '平均心率,最慢心率,最快心率,腿动总次数REM,腿动总次数NREM,腿动总次数Sleep'
    sheet_head=sheet_head.split(',')

    # 写入excel标题栏
    title_count=0
    for title in sheet_head:
        sheet.write(0, title_count, title)
        title_count=title_count+1

    '''
    --------------------------------------------
    2. 打开所有文档，依次读取并写入数据
    --------------------------------------------
    '''
    patient_count=0
    file_wrong_num=0

    with open(group+'\\'+"文件待修改格式.txt", "w") as f:
        f.write('待修改格式文件清单\r\n')

    for file in os.listdir('.'+'\\'+group+'\\docx\\'):
        if file[-4:] == 'docx':
            try:
                # 打开文档
                document = Document('.'+'\\'+group+'\\docx\\'+file)
                patient_count=patient_count+1

                # 读取表格材料
                tables = [table for table in document.tables]
                last_cell=0

                # 个人信息
                table_info=tables[0]
                count=0
                for row in table_info.rows:
                    for cell in row.cells:
                        if count == 3:
                            sheet.write(patient_count, sheet_head.index('姓名'), cell.text.split(':')[1])
                        if count == 4:
                            sheet.write(patient_count, sheet_head.index('性别'), cell.text.split(':')[1])
                        if count == 5:
                            # 在此处更改需要计算年龄的年份。 需要修改两处！！
                            sheet.write(patient_count, sheet_head.index('Date of birth'), cell.text.split(':')[1].replace('-', '/'))
                            if len(cell.text.split(':')[1].split('-'))>1:
                                sheet.write(patient_count, sheet_head.index('年龄'), 2012-float(cell.text.split(':')[1].split('-')[0]))
                            if len(cell.text.split(':')[1].split('/'))>1:
                                sheet.write(patient_count, sheet_head.index('年龄'),
                                            2012 - float(cell.text.split(':')[1].split('/')[0]))
                        if count == 6:
                            sheet.write(patient_count, sheet_head.index('身高'), int(cell.text[cell.text.index('身高')+3:cell.text.index('cm')]))
                        if count == 7:
                            sheet.write(patient_count, sheet_head.index('体重'), cell.text[cell.text.index('体重')+3:cell.text.index('kg')])
                        if count == 8:
                            sheet.write(patient_count, sheet_head.index('BMI'), cell.text[cell.text.index('BMI')+4:cell.text.index('kg')])
                        if count == 9:
                            sheet.write(patient_count, sheet_head.index('Date of study'), cell.text.split(':')[1].replace('-', '/'))
                            sheet.write(patient_count, sheet_head.index('监测日期'), cell.text.split(':')[1].replace('-', '/').replace(' ', ''))
                        count = count+1
                #print('个人信息finished...')

                # 睡眠摘要
                table_sleep_abstract=tables[1]
                count=0
                # totaltime=1
                for row in table_sleep_abstract.rows:
                    for cell in row.cells:
                        if count == 9:
                            sheet.write(patient_count, sheet_head.index('总记录时间'), cell.text.split(':')[1])
                        if count == 11:
                            sheet.write(patient_count, sheet_head.index('有效睡眠时间'), cell.text.split(':')[1])
                        if count == 15:
                            sheet.write(patient_count, sheet_head.index('总睡眠时间'),
                                        float(cell.text.split(':')[1]) * 60 + float(cell.text.split(':')[2]))
                            totaltime=(float(cell.text.split(':')[1]) + float(cell.text.split(':')[2]) / 60 )
                        if count == 17:
                            sheet.write(patient_count, sheet_head.index('睡眠潜伏期'), cell.text.split(':')[1])
                        if count == 19:
                            sheet.write(patient_count, sheet_head.index('睡眠效率'), cell.text.split(':')[1])

                        if count == 30:
                            sheet.write(patient_count, sheet_head.index('Wake时间'), cell.text)
                        if count == 33:
                            sheet.write(patient_count, sheet_head.index('StageREM时间'), cell.text)
                        if count == 35:
                            sheet.write(patient_count, sheet_head.index('StageREM占比'), cell.text)
                        if count == 37:
                            sheet.write(patient_count, sheet_head.index('Stage1时间'), cell.text)
                        if count == 39:
                            sheet.write(patient_count, sheet_head.index('Stage1占比'), cell.text)
                        if count == 41:
                            sheet.write(patient_count, sheet_head.index('Stage2时间'), cell.text)
                        if count == 43:
                            sheet.write(patient_count, sheet_head.index('Stage2占比'), cell.text)
                        if count == 45:
                            sheet.write(patient_count, sheet_head.index('Stage3时间'), cell.text)
                        if count == 47:
                            sheet.write(patient_count, sheet_head.index('Stage3占比'), cell.text)
                       # if count == 49:
                       #     sheet.write(patient_count, sheet_head.index('Stage4时间'), cell.text)
                       # if count == 51:
                       #     sheet.write(patient_count, sheet_head.index('Stage4占比'), cell.text)
                       # if count == 56:
                       #     sheet.write(patient_count, sheet_head.index('觉醒次数'), cell.text.split(':')[1])
                       # if count == 58:
                       #     sheet.write(patient_count, sheet_head.index('觉醒指数'), cell.text.split(':')[1])
                        if count == 52:
                            sheet.write(patient_count, sheet_head.index('觉醒次数'), cell.text.split(':')[1])
                        if count == 54:
                            sheet.write(patient_count, sheet_head.index('觉醒指数'), cell.text.split(':')[1])
                            
                        count=count+1
	
                #print('睡眠摘要Finished...')

                # 呼吸事件统计
                table_breath=tables[2]
                count=0

                for row in table_breath.rows:
                    for cell in row.cells:
                        if count == 15:
                            sheet.write(patient_count, sheet_head.index('阻塞性事件次数'), cell.text)
                        if count == 16:
                            sheet.write(patient_count, sheet_head.index('混合性事件次数'), cell.text)
                        if count == 17:
                            sheet.write(patient_count, sheet_head.index('中枢性事件次数'), cell.text)
                        if count == 18:
                            times_apnea=int(cell.text)
                            sheet.write(patient_count, sheet_head.index('所有暂停次数'), cell.text)
                        if count == 19:
                            times_hypo=int(cell.text)
                            sheet.write(patient_count, sheet_head.index('低通气次数'), cell.text)

                        if count == 22:
                            sheet.write(patient_count, sheet_head.index('阻塞性事件平均时间'), cell.text)
                        if count == 23:
                            sheet.write(patient_count, sheet_head.index('混合性事件平均时间'), cell.text)
                        if count == 24:
                            sheet.write(patient_count, sheet_head.index('中枢性事件平均时间'), cell.text)
                        if count == 25:
                            meantime_apnea=float(cell.text)
                            sheet.write(patient_count, sheet_head.index('所有暂停平均时间'), cell.text)
                        if count == 26:
                            meantime_hypo=float(cell.text)
                            sheet.write(patient_count, sheet_head.index('低通气平均时间'), cell.text)
                        if count == 32:
                            sheet.write(patient_count, sheet_head.index('所有暂停最长时间'), cell.text)
                        if count == 33:
                            sheet.write(patient_count, sheet_head.index('低通气最长时间'), cell.text)

                        if count == 36:
                            sheet.write(patient_count, sheet_head.index('阻塞性事件NREM下次数'), cell.text)
                        if count == 37:
                            sheet.write(patient_count, sheet_head.index('混合性事件NREM下次数'), cell.text)
                        if count == 38:
                            sheet.write(patient_count, sheet_head.index('中枢性事件NREM下次数'), cell.text)
                        if count == 39:
                            sheet.write(patient_count, sheet_head.index('所有暂停NREM下次数'), cell.text)
                        if count == 40:
                            sheet.write(patient_count, sheet_head.index('低通气NREM下次数'), cell.text)

                        if count == 43:
                            sheet.write(patient_count, sheet_head.index('阻塞性事件REM下次数'), cell.text)
                        if count == 44:
                            sheet.write(patient_count, sheet_head.index('混合性事件REM下次数'), cell.text)
                        if count == 45:
                            sheet.write(patient_count, sheet_head.index('中枢性事件REM下次数'), cell.text)
                        if count == 46:
                            sheet.write(patient_count, sheet_head.index('所有暂停REM下次数'), cell.text)
                        if count == 47:
                            sheet.write(patient_count, sheet_head.index('低通气REM下次数'), cell.text)
                        if count == 49:
                            sheet.write(patient_count, sheet_head.index('AHI'), cell.text.split('：')[1].split('O')[0])
                            sheet.write(patient_count, sheet_head.index('OAI'), cell.text.split('：')[2])
                        if count == 56:
                            sheet.write(patient_count, sheet_head.index('鼾声次数'), cell.text[cell.text.index('鼾声次数')+5:cell.text.index('鼾声指数')])
                            sheet.write(patient_count, sheet_head.index('鼾声指数'), cell.text.split('：')[2])
                        count = count+1
                sheet.write(patient_count, sheet_head.index('MAD'), (meantime_hypo*times_hypo+times_apnea*meantime_apnea)/(times_hypo+times_apnea))
                #print('呼吸事件统计Finished...')

                # 呼吸体位
                table_breath_position=tables[3]
                count=0
                for col in table_breath_position.columns:
                    for cell in col.cells:
                        if count ==9:
                            sheet.write(patient_count, sheet_head.index('俯卧下阻塞性指数'), cell.text)
                        if count == 10:
                            sheet.write(patient_count, sheet_head.index('左侧卧下阻塞性指数'), cell.text)
                        if count == 11:
                            sheet.write(patient_count, sheet_head.index('右侧卧下阻塞性指数'), cell.text)
                        if count == 12:
                            sheet.write(patient_count, sheet_head.index('仰卧下阻塞性指数'), cell.text)

                        if count == 16:
                            sheet.write(patient_count, sheet_head.index('俯卧下混合性指数'), cell.text)
                        if count == 17:
                            sheet.write(patient_count, sheet_head.index('左侧卧下混合性指数'), cell.text)
                        if count == 18:
                            sheet.write(patient_count, sheet_head.index('右侧卧下混合性指数'), cell.text)
                        if count == 19:
                            sheet.write(patient_count, sheet_head.index('仰卧下混合性指数'), cell.text)

                        if count == 23:
                            sheet.write(patient_count, sheet_head.index('俯卧下中枢性指数'), cell.text)
                        if count == 24:
                            sheet.write(patient_count, sheet_head.index('左侧卧下中枢性指数'), cell.text)
                        if count == 25:
                            sheet.write(patient_count, sheet_head.index('右侧卧下中枢性指数'), cell.text)
                        if count == 26:
                            sheet.write(patient_count, sheet_head.index('仰卧下中枢性指数'), cell.text)

                        if count ==30:
                            sheet.write(patient_count, sheet_head.index('俯卧下低通气指数'), cell.text)
                        if count == 31:
                            sheet.write(patient_count, sheet_head.index('左侧卧下低通气指数'), cell.text)
                        if count == 32:
                            sheet.write(patient_count, sheet_head.index('右侧卧下低通气指数'), cell.text)
                        if count == 33:
                            sheet.write(patient_count, sheet_head.index('仰卧下低通气指数'), cell.text)

                        count=count+1
                #print('呼吸体位统计Finished...')

                #血氧摘要
                table_oxygen=tables[4]
                count=0
                # odindexa=0
                # odindexb=0
                for col in table_oxygen.columns:
                    for cell in col.cells:
                        if count == 1:
                            sheet.write(patient_count, sheet_head.index('平均血氧'), cell.text.split('：')[1])
                        if count == 2:
                            sheet.write(patient_count, sheet_head.index('最低血氧'), cell.text.split('：')[1])
                        if count == 20:
                            sheet.write(patient_count, sheet_head.index('氧减值>=2%次数'), cell.text)
                        if count == 21:
                            sheet.write(patient_count, sheet_head.index('氧减值>=3%次数'), cell.text)
                        if count == 22:
                            sheet.write(patient_count, sheet_head.index('氧减值>=4%次数'), cell.text)
                            odindexa=(float(cell.text))
                            odindexb=(float(odindexa) / float(totaltime))
                            sheet.write(patient_count, sheet_head.index('ODI'), float(odindexb))
                        if count == 23:
                            sheet.write(patient_count, sheet_head.index('氧减值>=5%次数'), cell.text)

                        if count == 39:
                            sheet.write(patient_count, sheet_head.index('平均心率'), cell.text)
                        if count == 40:
                            sheet.write(patient_count, sheet_head.index('最慢心率'), cell.text)
                        if count == 41:
                            sheet.write(patient_count, sheet_head.index('最快心率'), cell.text)

                        if count == 48:
                            sheet.write(patient_count, sheet_head.index('血氧<95%时间（分钟）'), float(cell.text.split(':')[0])*60+float(cell.text.split(':')[1])+float(cell.text.split(':')[2])/60)
                        if count == 49:
                            sheet.write(patient_count, sheet_head.index('血氧<90%时间（分钟）'), float(cell.text.split(':')[0])*60+float(cell.text.split(':')[1])+float(cell.text.split(':')[2])/60)
                        if count == 50:
                            sheet.write(patient_count, sheet_head.index('血氧<85%时间（分钟）'), float(cell.text.split(':')[0])*60+float(cell.text.split(':')[1])+float(cell.text.split(':')[2])/60)
                        if count == 51:
                            sheet.write(patient_count, sheet_head.index('血氧<80%时间（分钟）'), float(cell.text.split(':')[0])*60+float(cell.text.split(':')[1])+float(cell.text.split(':')[2])/60)
                        count = count+1
                #print('血氧摘要Finished...')

                # 腿动统计
                table_leg=tables[5]
                count=0
                for row in table_leg.rows:
                    for cell in row.cells:
                        if count == 9:
                            sheet.write(patient_count, sheet_head.index('腿动总次数REM'), cell.text)
                        if count == 10:
                            sheet.write(patient_count, sheet_head.index('腿动总次数NREM'), cell.text)
                        if count == 11:
                            sheet.write(patient_count, sheet_head.index('腿动总次数Sleep'), cell.text)
                        count=count+1
                # print('腿动统计Finished...')
                print('正统计第' + str(patient_count) + '名患者...')
            except Exception as e:
                file_wrong_num=file_wrong_num+1
                print('文件'+file+'格式错误，请手动统计')
                print (traceback.format_exc())

                with open(group+'\\'+"文件待修改格式.txt", "a") as f:
                    f.write('文件' + file + '格式错误，请手动统计')
                    f.write(traceback.format_exc())
                    f.write('\n')

    print('共统计'+str(patient_count)+'名患者...')
    print('写入成功'+str(patient_count-file_wrong_num)+'写入失败'+str(file_wrong_num))
    print("写入数据成功！")

    with open(group+'\\'+"文件待修改格式.txt", "a") as f:
        f.write('共统计'+str(patient_count)+'名患者...')
        f.write('写入成功'+str(patient_count-file_wrong_num)+'写入失败'+str(file_wrong_num))

workbook.save('result.xls')  # 保存工作簿
